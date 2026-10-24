"""Extrakcia textu a metadát z HTML, PDF a DOCX + zistenie jazyka."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import py3langid
import trafilatura

from .textnorm import clean_ws

# Jazyky, ktoré ATENA podporuje; detektor sa obmedzí na ne (presnejšie výsledky)
SUPPORTED = ["sk", "cs", "en", "de", "pl", "hu", "ru", "zh", "fr", "fi", "sv", "nl", "it", "es", "uk", "et", "ja", "ko"]
_identifier = py3langid.langid.LanguageIdentifier.from_modelpath(
    py3langid.langid.MODEL_DIR / py3langid.langid.MODEL_FILE, norm_probs=True)
_identifier.set_languages(SUPPORTED)


@dataclass
class Extracted:
    title: str = ""
    text: str = ""
    language: str | None = None
    published: str | None = None
    author: str | None = None
    mime: str = ""
    pages: int | None = None
    meta: dict = field(default_factory=dict)


def detect_language(text: str) -> str | None:
    sample = text[:4000]
    if len(sample.strip()) < 40:
        return None
    lang, prob = _identifier.classify(sample)
    # slovenčina a čeština sú si blízke – rozhodne pri neistote pomer typických znakov
    if lang in ("sk", "cs"):
        sk = len(re.findall(r"[ľĺŕäô]", sample, re.I))
        cs = len(re.findall(r"[řůě]", sample, re.I))
        if sk > cs * 1.5 and sk >= 2:
            lang = "sk"
        elif cs > sk * 1.5 and cs >= 2:
            lang = "cs"
    return lang if prob >= 0.5 else None


def from_html(content: bytes, url: str) -> Extracted:
    html = content.decode("utf-8", errors="replace") if isinstance(content, bytes) else content
    doc = trafilatura.bare_extraction(html, url=url, include_comments=False, include_tables=True,
                                      favor_precision=True, with_metadata=True)
    text, title, date, author = "", "", None, None
    if doc is not None:
        d = doc.as_dict() if hasattr(doc, "as_dict") else dict(doc)
        text = d.get("text") or ""
        title = d.get("title") or ""
        date = d.get("date")
        author = d.get("author")
    if not title:
        m = re.search(r"<title[^>]*>(.*?)</title>", html, re.I | re.S)
        title = clean_ws(re.sub(r"<[^>]+>", "", m.group(1))) if m else ""
    lang = detect_language(text or title)
    m = re.search(r"<html[^>]+lang=[\"']?([a-zA-Z]{2})", html)
    if not lang and m:
        lang = m.group(1).lower()
    return Extracted(title=clean_ws(title), text=text.strip(), language=lang, published=date, author=author, mime="text/html")


def from_pdf(content: bytes) -> Extracted:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    parts = []
    for i, page in enumerate(reader.pages):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        parts.append(f"[strana {i + 1}]\n{t.strip()}")   # číslo strany kvôli citáciám
    text = "\n\n".join(parts)
    meta = reader.metadata or {}
    title = clean_ws(str(meta.get("/Title") or "")) if meta else ""
    if not title:
        first = next((ln for ln in text.splitlines() if len(ln.strip()) > 8 and not ln.startswith("[strana")), "")
        title = clean_ws(first)[:200]
    ex = Extracted(title=title, text=text, language=detect_language(text), mime="application/pdf", pages=len(reader.pages))
    # málo textu = pravdepodobne sken; OCR (Tesseract) je samostatný voliteľný krok
    ex.meta["needs_ocr"] = len(re.sub(r"\[strana \d+\]|\s", "", text)) < 50 * max(1, len(reader.pages))
    return ex


def from_docx(content: bytes) -> Extracted:
    import docx
    d = docx.Document(io.BytesIO(content))
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            paras.append(" | ".join(c.text.strip() for c in row.cells))
    text = "\n".join(paras)
    title = d.core_properties.title or (paras[0][:200] if paras else "")
    return Extracted(title=clean_ws(title), text=text, language=detect_language(text),
                     mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     author=d.core_properties.author or None)


def extract(content: bytes, content_type: str, url: str) -> Extracted:
    ct = (content_type or "").lower()
    path = url.lower().split("?")[0]
    if "pdf" in ct or path.endswith(".pdf") or content[:5] == b"%PDF-":
        return from_pdf(content)
    if "wordprocessingml" in ct or path.endswith(".docx"):
        return from_docx(content)
    if ct.startswith("text/plain"):
        text = content.decode("utf-8", errors="replace")
        return Extracted(title=text.strip().splitlines()[0][:200] if text.strip() else "", text=text,
                         language=detect_language(text), mime="text/plain")
    return from_html(content, url)
